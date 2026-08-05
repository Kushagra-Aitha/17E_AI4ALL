"""Streamlit demo for the emergency department wait-time and admission models."""

from __future__ import annotations

import csv
from collections import Counter
from io import BytesIO, StringIO
import importlib
import json
from pathlib import Path

import joblib
import numpy as np
import pandas as pd
import streamlit as st

try:
    from streamlit_app import wait_time_tab
except ModuleNotFoundError:
    import wait_time_tab

try:
    from streamlit_app import colors
except ModuleNotFoundError:
    import colors

wait_time_tab = importlib.reload(wait_time_tab)
render_wait_time_tab = wait_time_tab.render_wait_time_tab


ROOT = Path(__file__).resolve().parent.parent
MODELS_DIR = ROOT / "models"
ASSETS_DIR = Path(__file__).resolve().parent / "assets"
YALE_UPLOAD_COLUMNS = [
    "esi",
    "age",
    "gender",
    "race",
    "ethnicity",
    "insurance_status",
    "arrivalmode",
    "arrivalmonth",
    "arrivalday",
    "arrivalhour_bin",
    "triage_vital_hr",
    "triage_vital_sbp",
    "triage_vital_dbp",
    "triage_vital_rr",
    "triage_vital_o2",
    "triage_vital_o2_device",
    "triage_vital_temp",
    "n_edvisits",
    "n_admissions",
    "chief_complaints",
]
YALE_COMPLAINT_LABELS = {
    "cc_abdominalpain": "abdominal pain",
    "cc_alcoholintoxication": "alcohol intoxication",
    "cc_breathingdifficulty": "breathing difficulty",
    "cc_cellulitis": "cellulitis",
    "cc_chestpain": "chest pain",
    "cc_dentalpain": "dental pain",
    "cc_fall_65": "fall, age 65 or older",
    "cc_fever": "fever",
    "cc_fulltrauma": "full trauma",
    "cc_gibleeding": "GI bleeding",
    "cc_headache_newonsetornewsymptoms": "headache, new onset or new symptoms",
    "cc_headinjury": "head injury",
    "cc_laceration": "laceration",
    "cc_motorvehiclecrash": "motor vehicle crash",
    "cc_respiratorydistress": "respiratory distress",
    "cc_shortnessofbreath": "shortness of breath",
    "cc_strokealert": "stroke alert",
    "cc_suture_stapleremoval": "suture/staple removal",
}


def load_json(filename: str):
    return json.loads((MODELS_DIR / filename).read_text(encoding="utf-8"))


@st.cache_resource
def load_artifacts():
    return {
        "nhamcs_model1": joblib.load(MODELS_DIR / "nhamcs_model1.joblib"),
        "nhamcs_model2": joblib.load(MODELS_DIR / "nhamcs_model2.joblib"),
        "nhamcs_features1": load_json("nhamcs_model1_features.json"),
        "nhamcs_features2": load_json("nhamcs_model2_features.json"),
        "nhamcs_sample1": load_json("sample_input_model1.json"),
        "nhamcs_sample2": load_json("sample_input_model2.json"),
        "nhamcs_metrics": load_json("nhamcs_metrics.json"),
        "yale_model": joblib.load(MODELS_DIR / "yale_admission_model.joblib"),
        "yale_features": load_json("yale_features.json"),
        "yale_metrics": load_json("yale_metrics.json"),
    }


def keep_tab_selected(tab_name: str) -> None:
    st.session_state["active_tab"] = tab_name


def go_to_wait_time(mode: str) -> None:
    keep_tab_selected("Wait-Time Prediction")
    st.session_state["wait_input_mode"] = mode


def go_to_admission() -> None:
    keep_tab_selected("Admission Risk")


def set_one_hot(row: dict, prefix: str, selection: str) -> None:
    columns = [column for column in row if column.startswith(prefix)]
    for column in columns:
        row[column] = 0
    selected_column = f"{prefix}{selection}"
    if selected_column in row:
        row[selected_column] = 1


def make_nhamcs_row(
    sample: dict,
    age: int,
    payment_type: str,
    arrival_mode: str,
    triage_acuity: str,
    pulse: int,
    systolic_bp: int,
) -> dict:
    row = sample.copy()
    row["age_years"] = age
    set_one_hot(row, "payment_type_", payment_type)
    set_one_hot(row, "arrival_mode_", arrival_mode)
    set_one_hot(row, "triage_acuity_", triage_acuity)

    if "pulse" in row:
        row["pulse"] = pulse
        row["pulse_missing"] = 0
    if "systolic_bp" in row:
        row["systolic_bp"] = systolic_bp
        row["systolic_bp_missing"] = 0
    return row


def predict_wait_minutes(bundle: dict, features: list[str], row: dict) -> float:
    if bundle.get("target") != "log_wait_time" or bundle.get("target_transform") != "log1p":
        raise ValueError("Unexpected NHAMCS target format; expected log1p wait time.")

    frame = pd.DataFrame([row], columns=features)
    numeric_columns = bundle["numeric_columns"]
    frame[numeric_columns] = bundle["scaler"].transform(frame[numeric_columns])
    log_wait = float(bundle["model"].predict(frame)[0])
    return max(0.0, float(np.expm1(log_wait)))


def duplicate_headers(headers: list) -> list[str]:
    names = [str(header) for header in headers if header not in (None, "")]
    counts = Counter(names)
    return sorted(name for name, count in counts.items() if count > 1)


def read_yale_upload(file_name: str, content: bytes) -> pd.DataFrame:
    if not content:
        raise ValueError("The uploaded file is empty.")

    lower_name = file_name.lower()
    if lower_name.endswith(".csv"):
        try:
            text = content.decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise ValueError("CSV files must use UTF-8 encoding.") from exc
        reader = csv.reader(StringIO(text))
        try:
            headers = next(reader)
        except StopIteration as exc:
            raise ValueError("The uploaded file is empty.") from exc
        duplicates = duplicate_headers(headers)
        if duplicates:
            raise ValueError(f"Duplicate column names found: {', '.join(duplicates)}")
        try:
            return pd.read_csv(BytesIO(content))
        except pd.errors.EmptyDataError as exc:
            raise ValueError("The uploaded file is empty.") from exc

    if lower_name.endswith(".xlsx"):
        from openpyxl import load_workbook

        try:
            workbook = load_workbook(
                BytesIO(content),
                read_only=True,
                data_only=True,
            )
            worksheet = workbook.active
            header_row = next(
                worksheet.iter_rows(min_row=1, max_row=1, values_only=True),
                None,
            )
            workbook.close()
        except Exception as exc:
            raise ValueError("The Excel file could not be read.") from exc
        if header_row is None:
            raise ValueError("The uploaded file is empty.")
        duplicates = duplicate_headers(list(header_row))
        if duplicates:
            raise ValueError(f"Duplicate column names found: {', '.join(duplicates)}")
        return pd.read_excel(BytesIO(content))

    raise ValueError("Upload a .csv or .xlsx file.")


def yale_category_values(model, feature_info: dict) -> dict[str, set]:
    categorical_pipeline = model.named_steps["preprocessor"].named_transformers_[
        "categorical"
    ]
    onehot = categorical_pipeline.named_steps["onehot"]
    return {
        column: set(categories.tolist())
        for column, categories in zip(
            feature_info["categorical_columns"],
            onehot.categories_,
        )
    }


def format_spreadsheet_rows(mask: pd.Series) -> str:
    row_numbers = (mask[mask].index + 2).tolist()
    displayed = ", ".join(str(row) for row in row_numbers[:10])
    return f"{displayed}..." if len(row_numbers) > 10 else displayed


def prepare_yale_upload(
    uploaded_data: pd.DataFrame,
    model,
    feature_info: dict,
) -> pd.DataFrame:
    if uploaded_data.columns.duplicated().any():
        duplicates = uploaded_data.columns[uploaded_data.columns.duplicated()].tolist()
        raise ValueError(f"Duplicate column names found: {', '.join(duplicates)}")

    missing_columns = [
        column for column in YALE_UPLOAD_COLUMNS if column not in uploaded_data.columns
    ]
    if missing_columns:
        raise ValueError(f"Missing required columns: {', '.join(missing_columns)}")
    if uploaded_data.empty:
        raise ValueError("The uploaded file has no data rows.")

    data = uploaded_data.reset_index(drop=True).copy()
    errors = []
    numeric_values = pd.DataFrame(index=data.index)

    for column in feature_info["numeric_columns"]:
        values = pd.to_numeric(data[column], errors="coerce")
        invalid = values.isna() | ~np.isfinite(values)
        if invalid.any():
            errors.append(
                f"{column} must contain a finite numeric value "
                f"(rows {format_spreadsheet_rows(invalid)})"
            )
        numeric_values[column] = values

    invalid_esi = ~numeric_values["esi"].isin([1, 2, 3, 4, 5])
    if invalid_esi.any():
        errors.append(
            "esi must be one of 1, 2, 3, 4, or 5 "
            f"(rows {format_spreadsheet_rows(invalid_esi)})"
        )

    invalid_device = ~numeric_values["triage_vital_o2_device"].isin([0, 1])
    if invalid_device.any():
        errors.append(
            "triage_vital_o2_device must be 0 or 1 "
            f"(rows {format_spreadsheet_rows(invalid_device)})"
        )

    for column in ("n_edvisits", "n_admissions"):
        values = numeric_values[column]
        invalid_count = (values < 0) | (values % 1 != 0)
        if invalid_count.any():
            errors.append(
                f"{column} must be a non-negative whole number "
                f"(rows {format_spreadsheet_rows(invalid_count)})"
            )

    category_values = yale_category_values(model, feature_info)
    for column, valid_values in category_values.items():
        invalid_category = data[column].isna() | ~data[column].isin(valid_values)
        if invalid_category.any():
            errors.append(
                f"{column} contains an unknown category "
                f"(rows {format_spreadsheet_rows(invalid_category)}). "
                f"Allowed values: {', '.join(sorted(str(value) for value in valid_values))}"
            )

    valid_complaints = set(feature_info["cc_columns"])
    parsed_complaints = []
    for index, value in data["chief_complaints"].items():
        if pd.isna(value) or not str(value).strip():
            errors.append(f"chief_complaints is empty (row {index + 2})")
            parsed_complaints.append([])
            continue
        complaints = [item.strip() for item in str(value).split(";") if item.strip()]
        unknown = sorted(set(complaints) - valid_complaints)
        if unknown:
            errors.append(
                f"chief_complaints contains unknown code(s) at row {index + 2}: "
                f"{', '.join(unknown)}"
            )
        parsed_complaints.append(complaints)

    if errors:
        raise ValueError("\n".join(errors))

    model_input = pd.DataFrame(
        0,
        index=data.index,
        columns=feature_info["features"],
    )
    for column in feature_info["numeric_columns"]:
        model_input[column] = numeric_values[column]
    for column in feature_info["categorical_columns"]:
        model_input[column] = data[column]
    for index, complaints in enumerate(parsed_complaints):
        for complaint in complaints:
            model_input.at[index, complaint] = 1
    return model_input[feature_info["features"]]


def predict_yale_upload(
    uploaded_data: pd.DataFrame,
    model,
    feature_info: dict,
) -> pd.DataFrame:
    model_input = prepare_yale_upload(uploaded_data, model, feature_info)
    probabilities = model.predict_proba(model_input)[:, 1]
    predictions = model.predict(model_input).astype(int)

    results = uploaded_data.reset_index(drop=True).copy()
    results["admission_probability"] = probabilities
    results["predicted_disposition"] = np.where(predictions == 1, "Admit", "Discharge")
    return results


def yale_sample_upload() -> pd.DataFrame:
    rows = [
        [2, 67, "Male", "White or Caucasian", "Non-Hispanic", "Medicare", "ambulance", "July", "Tuesday", "19-22", 112, 158, 92, 22, 94, 0, 98.7, 2, 1, "cc_chestpain"],
        [1, 74, "Female", "Black or African American", "Non-Hispanic", "Medicaid", "ambulance", "January", "Sunday", "23-02", 128, 92, 58, 32, 86, 1, 103.1, 3, 1, "cc_respiratorydistress;cc_fever"],
        [2, 58, "Male", "Asian", "Non-Hispanic", "Commercial", "Car", "October", "Friday", "07-10", 118, 101, 64, 24, 95, 0, 98.4, 1, 0, "cc_gibleeding"],
        [1, 81, "Female", "White or Caucasian", "Non-Hispanic", "Medicare", "ambulance", "March", "Monday", "11-14", 88, 190, 105, 20, 93, 0, 99.0, 4, 2, "cc_strokealert"],
        [3, 49, "Male", "Black or African American", "Hispanic or Latino", "Medicaid", "Walk-in", "August", "Wednesday", "15-18", 102, 132, 78, 18, 97, 0, 100.8, 2, 0, "cc_cellulitis"],
        [4, 29, "Female", "Asian", "Non-Hispanic", "Commercial", "Car", "May", "Saturday", "19-22", 84, 124, 76, 16, 99, 0, 98.2, 0, 0, "cc_laceration"],
        [5, 35, "Male", "Other", "Hispanic or Latino", "Self pay", "Public Transportation", "June", "Thursday", "07-10", 78, 126, 82, 16, 99, 0, 98.5, 1, 0, "cc_dentalpain"],
        [4, 44, "Male", "White or Caucasian", "Non-Hispanic", "Medicaid", "Police", "December", "Saturday", "23-02", 96, 138, 86, 18, 97, 0, 97.9, 5, 0, "cc_alcoholintoxication"],
        [3, 53, "Female", "American Indian or Alaska Native", "Non-Hispanic", "Commercial", "Car", "September", "Monday", "03-06", 108, 116, 72, 20, 96, 0, 99.6, 1, 0, "cc_abdominalpain"],
        [5, 41, "Female", "Native Hawaiian or Other Pacific Islander", "Non-Hispanic", "Commercial", "Walk-in", "February", "Wednesday", "11-14", 72, 120, 74, 14, 99, 0, 98.1, 0, 0, "cc_suture_stapleremoval"],
    ]
    return pd.DataFrame(rows, columns=YALE_UPLOAD_COLUMNS)


def yale_prediction_summary(results: pd.DataFrame) -> dict:
    disposition_counts = results["predicted_disposition"].value_counts()
    return {
        "total": len(results),
        "admit": int(disposition_counts.get("Admit", 0)),
        "discharge": int(disposition_counts.get("Discharge", 0)),
        "highest_probability": float(results["admission_probability"].max()),
        "lowest_probability": float(results["admission_probability"].min()),
    }


def readable_complaint(code: str) -> str:
    if code in YALE_COMPLAINT_LABELS:
        return YALE_COMPLAINT_LABELS[code]
    return code.removeprefix("cc_").replace("_", " ")


def readable_complaints(value: str) -> str:
    codes = [code.strip() for code in str(value).split(";") if code.strip()]
    return " + ".join(readable_complaint(code) for code in codes)


def display_number(value) -> str:
    numeric_value = float(value)
    return str(int(numeric_value)) if numeric_value.is_integer() else f"{numeric_value:g}"


def admission_risk_status(probability: float) -> str:
    if probability >= 0.7:
        return "serious"
    if probability >= 0.3:
        return "warning"
    return "good"


_ADMISSION_RISK_LABELS = {"good": "Low", "warning": "Moderate", "serious": "High"}


def admission_risk_level(probability: float) -> str:
    return _ADMISSION_RISK_LABELS[admission_risk_status(probability)]


def admission_risk_badge_html(probability: float) -> str:
    status = admission_risk_status(probability)
    style = colors.STATUS_STYLE[status]
    level = _ADMISSION_RISK_LABELS[status]
    return (
        f'<span class="status-badge" style="background:{style["bg"]}; '
        f'color:{style["fg"]};">{style["icon"]} {level} chance of admission</span>'
    )


def yale_patient_label(row: pd.Series, row_number: int) -> str:
    if "patient_name" in row.index and pd.notna(row["patient_name"]):
        patient_name = str(row["patient_name"]).strip()
        if patient_name:
            return patient_name
    return f"Row {row_number}"


def yale_visit_summary(row: pd.Series) -> str:
    return (
        f"{display_number(row['age'])}-year-old, "
        f"ESI {display_number(row['esi'])}, "
        f"{readable_complaints(row['chief_complaints'])}"
    )


def yale_results_table(results: pd.DataFrame) -> pd.DataFrame:
    rows = results.reset_index(drop=True)
    risk_status = rows["admission_probability"].map(admission_risk_status)
    risk_icon = risk_status.map(lambda status: colors.STATUS_STYLE[str(status)]["icon"])
    risk_level = rows["admission_probability"].map(admission_risk_level)
    return pd.DataFrame(
        {
            "Patient": [
                yale_patient_label(row, index + 1)
                for index, row in rows.iterrows()
            ],
            "Chance of Admission": rows["admission_probability"].map(
                lambda probability: f"{probability:.1%}"
            ),
            "Risk": risk_icon + " " + risk_level,
            "Likely Outcome": rows["predicted_disposition"],
            "Summary": [yale_visit_summary(row) for _, row in rows.iterrows()],
        }
    )


def yale_risk_counts(results: pd.DataFrame) -> dict[str, int]:
    levels = results["admission_probability"].map(admission_risk_level)
    return {
        level: int((levels == level).sum())
        for level in ("High", "Moderate", "Low")
    }


def show_yale_visit_card(row: pd.Series, row_number: int) -> None:
    patient = yale_patient_label(row, row_number)
    probability = float(row["admission_probability"])
    prediction = row["predicted_disposition"]
    with st.container(border=True):
        st.markdown(f"#### {patient}")
        st.metric("Chance of Admission", f"{probability:.1%}")
        st.markdown(admission_risk_badge_html(probability), unsafe_allow_html=True)
        st.write(f"**Likely Outcome:** {prediction}")
        st.write(yale_visit_summary(row))


def yale_metrics_table(metrics: dict) -> pd.DataFrame:
    test = metrics["test"]
    return pd.DataFrame(
        [
            {
                "Accuracy": test["accuracy"],
                "Recall": test["recall"],
                "Precision": test["precision"],
                "F1": test["f1"],
                "ROC-AUC": metrics["test_roc_auc"],
            }
        ]
    )


_WAIT_MONITORING_CHECKS = [
    (
        "Prediction quality",
        "Track MAE and RMSE over time and inspect residuals for systematic "
        "under- or over-prediction.",
    ),
    (
        "Subgroup review",
        "Compare errors by race, ethnicity, insurance, arrival mode, and "
        "triage priority.",
    ),
    (
        "Data drift",
        "Watch for changes across years, hospitals, patient mix, coding, and "
        "missingness relative to the 2007-2022 NHAMCS data.",
    ),
    (
        "Operational gaps",
        "Staffing, available beds, queue length, crowding, and hospital "
        "capacity are not included and should be monitored separately.",
    ),
]

_ADMISSION_MONITORING_CHECKS = [
    (
        "Calibration",
        "Compare predicted admission chances with observed admission rates "
        "to determine whether the scores remain well calibrated.",
    ),
    (
        "Performance",
        "Track ROC-AUC, precision, recall, F1, and the confusion matrix over "
        "time. Review the prediction threshold and false negatives among "
        "patients who were admitted.",
    ),
    (
        "Fairness and validation",
        "Compare performance across demographic groups and complete "
        "external validation using data from other hospitals.",
    ),
    (
        "Data quality and drift",
        "Monitor chief complaint coding, missingness, ESI and vital-sign "
        "collection, and changes in triage or admission patterns.",
    ),
]


def render_monitoring_checks(checks: list[tuple[str, str]]) -> None:
    for title, detail in checks:
        with st.expander(title):
            st.write(detail)


LOGO_PATH = ASSETS_DIR / "logo.jpeg"

st.set_page_config(
    page_title="ED Model Explorer",
    page_icon=str(LOGO_PATH) if LOGO_PATH.exists() else "+",
    layout="wide",
)
if LOGO_PATH.exists():
    st.logo(str(LOGO_PATH), size="large")
st.markdown(
    """
    <style>
    .stApp p, .stApp li, .stApp label {
        font-size: 1.04rem;
        line-height: 1.55;
    }
    [data-testid="stCaptionContainer"] p {
        color: #4b5563;
        font-size: 0.96rem;
        line-height: 1.45;
    }
    [data-testid="stMetricLabel"] p {
        font-size: 0.98rem;
        font-weight: 600;
    }
    [data-testid="stMetricValue"] {
        font-size: 2rem;
    }
    [data-testid="stDataFrame"] {
        font-size: 1rem;
    }
    [data-testid="stFileUploader"] small {
        font-size: 0.9rem;
    }
    .status-badge {
        display: inline-flex;
        align-items: center;
        gap: 0.3rem;
        padding: 0.22rem 0.7rem;
        border-radius: 999px;
        font-size: 0.85rem;
        font-weight: 600;
        white-space: nowrap;
    }
    </style>
    """,
    unsafe_allow_html=True,
)
title_logo_col, title_text_col = st.columns([1, 11], vertical_alignment="center")
if LOGO_PATH.exists():
    with title_logo_col:
        st.image(str(LOGO_PATH), width="stretch")
with title_text_col:
    st.title("Drivers of Emergency Department Wait Times")


try:
    artifacts = load_artifacts()
except Exception as exc:
    st.error(f"Model artifacts could not be loaded: {exc}")
    st.stop()

tab_names = ["Home", "Wait-Time Prediction", "Admission Risk", "Monitoring"]
active_tab = st.session_state.get("active_tab", "Home")

with st.sidebar:
    
   

    st.markdown("**Jump to**")
    for name in tab_names:
        st.button(
            name,
            key=f"nav_{name}",
            width="stretch",
            type="primary" if name == active_tab else "secondary",
            disabled=name == active_tab,
            on_click=keep_tab_selected,
            args=(name,),
        )


if active_tab == "Home":

    st.write(
        "This app explores two emergency department outcomes: how long a patient may wait "
        "before first provider contact and the chance that a patient may be admitted after "
        "triage. Each workflow uses a separate dataset and model."
    )

    st.markdown("### Quick summary")
    nhamcs_display = artifacts["nhamcs_metrics"]["display"]
    yale_test = artifacts["yale_metrics"]["test"]
    summary_left, summary_right = st.columns(2)
    with summary_left:
        with st.container(border=True):
            st.markdown("**Wait-Time Prediction**")
            st.caption("Typical error when estimating minutes to first provider contact.")
            wait_metric_columns = st.columns(2)
            wait_metric_columns[0].metric(
                "R-squared", f"{nhamcs_display['model2']['r2']:.2f}"
            )
            wait_metric_columns[1].metric(
                "MAE", f"{nhamcs_display['model2']['mae_minutes']:.0f} min"
            )
    with summary_right:
        with st.container(border=True):
            st.markdown("**Admission Chance Screening**")
            st.caption("How well the model separates admitted from discharged visits.")
            yale_metric_columns = st.columns(2)
            yale_metric_columns[0].metric("Accuracy", f"{yale_test['accuracy']:.0%}")
            yale_metric_columns[1].metric(
                "ROC-AUC", f"{artifacts['yale_metrics']['test_roc_auc']:.2f}"
            )

    st.warning(
        "These outputs are for analysis, education, and review support, not clinical "
        "decisions. Wait-time estimates are approximate because staffing, bed availability, "
        "queue length, and crowding are not included. Admission estimates would require "
        "calibration, fairness checks, governance review, and external validation before "
        "real-world use."
    )

elif active_tab == "Wait-Time Prediction":
    with st.container(border=True):
        st.markdown("### Wait-Time Prediction")
        st.write(
            "Uses NHAMCS emergency department data to estimate approximate minutes before "
            "first provider contact. The model considers demographics, access factors, "
            "arrival mode, triage priority, and selected vital signs."
        )
        wait_choice_columns = st.columns(2)
        wait_choice_columns[0].button(
            "I have one patient",
            key="home_wait_single",
            width="stretch",
            on_click=go_to_wait_time,
            args=("Single patient form",),
        )
        wait_choice_columns[1].button(
            "I have a spreadsheet",
            key="home_wait_batch",
            width="stretch",
            on_click=go_to_wait_time,
            args=("CSV upload (many patients)",),
        )

    try:
        render_wait_time_tab(artifacts, keep_tab_selected)
    except Exception as exc:
        st.error("Wait-Time tab failed to load.")
        st.exception(exc)

elif active_tab == "Admission Risk":
    with st.container(border=True):
        st.markdown("### Admission Chance Screening")
        st.write(
            "Uses Yale ED triage data to estimate each visit's chance of admission. The "
            "upload workflow uses age, ESI, vitals, arrival details, insurance, prior "
            "utilization, and chief complaint codes."
        )
        admission_choice_columns = st.columns(2)
        admission_choice_columns[0].button(
            "I have one patient",
            key="home_admission_single",
            width="stretch",
            disabled=True,
            help="Single-patient entry isn't available yet — upload a one-row spreadsheet instead.",
        )
        admission_choice_columns[1].button(
            "I have a spreadsheet",
            key="home_admission_batch",
            width="stretch",
            on_click=go_to_admission,
        )

    st.subheader("Hospital Chance of Admission Screening")
    st.write(
        "Upload ED triage records and review the estimated chance of admission for each visit."
    )

    with st.container(border=True):
        yale_upload = st.file_uploader(
            "Upload CSV or Excel triage records",
            type=["csv", "xlsx"],
            key="yale_batch_upload",
            on_change=keep_tab_selected,
            args=("Admission Risk",),
        )
        with st.expander("Required upload format"):
            st.markdown("**Required columns**")
            st.code(", ".join(YALE_UPLOAD_COLUMNS), language=None)
            st.write(
                "Each row represents one ED visit. `chief_complaints` accepts one "
                "or more recognized `cc_*` codes separated by semicolons."
            )
            st.code(
                "cc_chestpain;cc_shortnessofbreath",
                language=None,
            )
            st.caption(
                "`patient_name` is optional and is used only as a display label. "
                "Full detailed results can be downloaded after prediction."
            )

    if yale_upload is not None:
        try:
            uploaded_yale_data = read_yale_upload(
                yale_upload.name,
                yale_upload.getvalue(),
            )
            upload_results = predict_yale_upload(
                uploaded_yale_data,
                artifacts["yale_model"],
                artifacts["yale_features"],
            )
        except (ValueError, TypeError, OSError) as exc:
            st.error(str(exc))
        else:
            st.success(f"Generated predictions for {len(upload_results):,} uploaded rows.")
            summary = yale_prediction_summary(upload_results)

            summary_columns = st.columns(5)
            summary_columns[0].metric("Total visits", f"{summary['total']:,}")
            summary_columns[1].metric("Likely admit", f"{summary['admit']:,}")
            summary_columns[2].metric(
                "Likely discharge",
                f"{summary['discharge']:,}",
            )
            summary_columns[3].metric(
                "Highest chance",
                f"{summary['highest_probability']:.1%}",
            )
            summary_columns[4].metric(
                "Lowest chance",
                f"{summary['lowest_probability']:.1%}",
            )

            st.markdown("**Risk mix in this upload**")
            risk_counts = yale_risk_counts(upload_results)
            risk_columns = st.columns(3)
            for risk_column, level in zip(risk_columns, ("Low", "Moderate", "High")):
                status = {"Low": "good", "Moderate": "warning", "High": "serious"}[level]
                style = colors.STATUS_STYLE[status]
                with risk_column:
                    st.markdown(
                        f'<div class="status-badge" style="background:{style["bg"]}; '
                        f'color:{style["fg"]}; width:100%; justify-content:center; '
                        f'font-size:1rem; padding:0.5rem;">'
                        f'{style["icon"]} {level}: {risk_counts[level]:,}</div>',
                        unsafe_allow_html=True,
                    )

            st.info(
                "These scores are for review support only and are not medical decisions."
            )

            ranked_results = upload_results.reset_index(drop=True).copy()
            ranked_results["_row_number"] = np.arange(1, len(ranked_results) + 1)

            st.markdown("### Top Cases in This Upload")
            top_visits = ranked_results.nlargest(3, "admission_probability")
            top_columns = st.columns(len(top_visits))
            for card_column, (_, row) in zip(top_columns, top_visits.iterrows()):
                with card_column:
                    show_yale_visit_card(row, int(row["_row_number"]))

            st.markdown("### Lowest Chance Cases")
            bottom_visits = ranked_results.nsmallest(3, "admission_probability")
            bottom_columns = st.columns(len(bottom_visits))
            for card_column, (_, row) in zip(bottom_columns, bottom_visits.iterrows()):
                with card_column:
                    show_yale_visit_card(row, int(row["_row_number"]))

            st.markdown("### All uploaded visits")
            st.dataframe(
                yale_results_table(upload_results),
                hide_index=True,
                use_container_width=True,
                column_config={
                    "Patient": st.column_config.TextColumn(width="medium"),
                    "Chance of Admission": st.column_config.TextColumn(width="small"),
                    "Risk": st.column_config.TextColumn(width="small"),
                    "Likely Outcome": st.column_config.TextColumn(width="small"),
                    "Summary": st.column_config.TextColumn(width="large"),
                },
            )
            st.download_button(
                "Download full prediction results",
                data=upload_results.to_csv(index=False).encode("utf-8"),
                file_name="yale_admission_predictions.csv",
                mime="text/csv",
                on_click=keep_tab_selected,
                args=("Admission Risk",),
            )

    with st.expander("Model performance details"):
        st.write(
            "The Yale model is balanced Logistic Regression trained on 560,484 ED triage "
            "visits. It uses 219 original predictors, expanded into 262 coefficient terms "
            "after preprocessing. Inputs cover demographics, insurance, arrival details, "
            "ESI, vital signs, prior utilization, and chief complaint indicators."
        )
        st.dataframe(
            yale_metrics_table(artifacts["yale_metrics"]),
            hide_index=True,
            width="stretch",
        )
        yale_test = artifacts["yale_metrics"]["test"]
        confusion_matrix = pd.DataFrame(
            yale_test["confusion_matrix"],
            index=["Actual discharge", "Actual admit"],
            columns=["Predicted discharge", "Predicted admit"],
        )
        st.markdown("**Yale test confusion matrix**")
        st.dataframe(confusion_matrix, width="stretch")

        yale_roc_plot = ASSETS_DIR / "yale_roc_curve.png"
        yale_esi_plot = ASSETS_DIR / "yale_admit_rate_by_esi.png"
        yale_plot_columns = st.columns(2, gap="medium")
        if yale_roc_plot.exists():
            with yale_plot_columns[0]:
                st.image(
                    yale_roc_plot,
                    caption=(
                        "Test ROC curve for the Yale admission model. ROC-AUC summarizes "
                        "overall separation between admitted and discharged visits."
                    ),
                    width="stretch",
                )
        if yale_esi_plot.exists():
            with yale_plot_columns[1]:
                st.image(
                    yale_esi_plot,
                    caption=(
                        "Observed admission rate by ESI level in the cleaned Yale dataset. "
                        "This is descriptive evidence, not a causal relationship."
                    ),
                    width="stretch",
                )

        st.info(
            "The Yale model separates admitted and discharged visits well overall, but real "
            "use would require calibration, subgroup fairness checks, governance review, "
            "and external validation."
        )

        st.markdown("**Supporting artifacts**")
        st.caption(
            "`models/yale_metrics.json`, feature definitions in `models/yale_features.json`, "
            "and the baseline Logistic Regression artifact."
        )

elif active_tab == "Monitoring":
    st.subheader("Monitoring Plan")
    st.write(
        "These checks describe what would need to be reviewed before and during real-world "
        "use. This page does not display live hospital monitoring."
    )
    st.caption("Click a check below to see what it means and why it matters.")

    wait_monitoring, admission_monitoring = st.columns(2)
    with wait_monitoring:
        with st.container(border=True):
            st.markdown("### Wait-Time Model Monitoring")
            render_monitoring_checks(_WAIT_MONITORING_CHECKS)


    with admission_monitoring:
        with st.container(border=True):
            st.markdown("### Admission Model Monitoring")
            render_monitoring_checks(_ADMISSION_MONITORING_CHECKS)

    with st.container(border=True):
        st.markdown("### Human Oversight")
        st.write(
            "These outputs support review only. Clinicians and hospital staff remain "
            "responsible for patient-care and operational decisions. Any real deployment "
            "would require documented intended use, privacy and governance review, local "
            "validation, periodic performance review, and clear retraining or rollback rules."
        )
